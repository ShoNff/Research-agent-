"""Word document generation tool."""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from claude_agent_sdk import tool


# Reliability tier to color mapping (RGB tuples)
TIER_COLORS = {
    "established": (34, 139, 34),   # Forest green
    "reputable": (70, 130, 180),     # Steel blue
    "emerging": (218, 165, 32),      # Goldenrod
    "opinion": (205, 92, 92),        # Indian red
    "unknown": (128, 128, 128),      # Gray
}


def _render_docx(report_data: dict, output_path: str, diagram_paths: list[str]) -> str:
    """Render a ReportDraft dict to a .docx file."""
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Style configuration
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)

    # Title
    title_para = doc.add_heading(report_data.get("title", "Research Report"), level=0)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Executive Summary
    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(report_data.get("executive_summary", ""))

    # Key Takeaways
    takeaways = report_data.get("key_takeaways", [])
    if takeaways:
        doc.add_heading("Key Takeaways", level=1)
        for takeaway in takeaways:
            doc.add_paragraph(takeaway, style="List Bullet")

    # Sections
    for section in report_data.get("sections", []):
        doc.add_heading(section.get("title", ""), level=2)
        doc.add_paragraph(section.get("content", ""))

        # Add source annotations for this section
        section_sources = section.get("sources", [])
        if section_sources:
            sources_para = doc.add_paragraph()
            run = sources_para.add_run("Sources: ")
            run.bold = True
            run.font.size = Pt(9)
            for i, src in enumerate(section_sources):
                tier = src.get("reliability_tier", "unknown")
                color = TIER_COLORS.get(tier, TIER_COLORS["unknown"])
                run = sources_para.add_run(
                    f"[{tier.upper()}] {src.get('title', 'Untitled')}"
                )
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(*color)
                if i < len(section_sources) - 1:
                    sources_para.add_run(" | ").font.size = Pt(9)

    # Embed diagrams
    diagram_path_list = diagram_paths or []
    for dp in diagram_path_list:
        p = Path(dp)
        if p.exists() and p.suffix in (".png", ".jpg", ".jpeg"):
            doc.add_heading("Diagram", level=2)
            doc.add_picture(str(p), width=Inches(6))

    # Full sources section
    all_sources = report_data.get("sources", [])
    if all_sources:
        doc.add_heading("References", level=1)
        # Group by tier
        by_tier: dict[str, list] = {}
        for src in all_sources:
            tier = src.get("reliability_tier", "unknown")
            by_tier.setdefault(tier, []).append(src)

        for tier_name in ["established", "reputable", "emerging", "opinion", "unknown"]:
            sources_in_tier = by_tier.get(tier_name, [])
            if not sources_in_tier:
                continue
            color = TIER_COLORS.get(tier_name, TIER_COLORS["unknown"])
            heading_para = doc.add_paragraph()
            run = heading_para.add_run(f"{tier_name.upper()} Sources")
            run.bold = True
            run.font.color.rgb = RGBColor(*color)

            for src in sources_in_tier:
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(src.get("title", "Untitled")).bold = True
                p.add_run(f"\n{src.get('url', '')}")

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    return str(output)


@tool(
    "render_docx",
    "Render a structured research report to a Word document (.docx). Pass the report as a JSON string matching the ReportDraft schema.",
    {
        "type": "object",
        "properties": {
            "report_json": {"type": "string", "description": "JSON-serialized ReportDraft"},
            "output_path": {"type": "string", "description": "Path for the output .docx file"},
            "diagram_paths": {
                "type": "array",
                "items": {"type": "string"},
                "description": "Paths to diagram images to embed",
            },
        },
        "required": ["report_json", "output_path"],
    },
)
async def render_docx(args: dict[str, Any]) -> dict[str, Any]:
    """Generate a Word document from report data."""
    try:
        report_data = json.loads(args["report_json"])
        diagram_paths = args.get("diagram_paths", [])
        path = _render_docx(report_data, args["output_path"], diagram_paths)
        return {
            "content": [{"type": "text", "text": f"Word document saved to: {path}"}]
        }
    except Exception as e:
        return {
            "content": [
                {"type": "text", "text": f"ERROR: Failed to generate Word document: {e}"}
            ]
        }
